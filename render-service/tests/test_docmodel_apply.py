# -*- coding: utf-8 -*-
"""V2 块级结构化模型：/parse 与 /apply（PRD §7.1）。

覆盖：
  1. parse audit-report：标题层级/段落/表格行列/占位符/tokens 快照断言；
  2. 原样回写：未编辑段落 XML 逐段一致、styles/numbering 部件字节不变、python-docx 可开；
  3. 改文字：段落/标题文本生效，其余段落 XML 不动，中文与 eastAsia 字体保留；
  4. 占位符：插入标量、删除占位符后 docxtpl 渲染正确；未配对标签报清晰错误；
  5. 表格：增行/删行/改单元格/标记循环行后 docxtpl 渲染出正确行数；
  6. unsupported：含 w:txbxContent 的段落 parse 标记 unsupported，apply 后 XML 原样保留；
  7. 安全：docm/宏、超大文件仍被 security.validate_docx 拒绝。
"""
import copy
import json
import zipfile
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from lxml import etree
from docx import Document
from docxtpl import DocxTemplate

from app.main import app
from app.docmodel import parse_docx
from app.docapply import apply_model
from app.preprocess import qn
from app.security import validate_docx, SecurityError

client = TestClient(app)

DOCX_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _blocks(resp_json):
    return resp_json["blocks"]


# --------------------------------------------------------------------------- #
# 1. parse audit-report 快照
# --------------------------------------------------------------------------- #
def test_parse_audit_report_structure(migrated_template_bytes):
    model = parse_docx(migrated_template_bytes)
    blocks = model["blocks"]

    # 顶层应有：控制段 p0、附录标题 p1、缺陷表 t0、若干控制段
    types = [(b["id"], b["type"]) for b in blocks]
    assert ("p0", "paragraph") in types
    assert ("p1", "heading") in types
    # p1 是附录编号段（style=9 outlineLvl=2 -> level 3）
    p1 = next(b for b in blocks if b["id"] == "p1")
    assert p1["level"] == 3

    # 表格 t0 存在，15 行
    t0 = next(b for b in blocks if b["id"] == "t0")
    assert t0["type"] == "table"
    assert len(t0["rows"]) == 15

    # p0 文本含 docxtpl 段落循环语法
    p0 = next(b for b in blocks if b["id"] == "p0")
    assert p0["text"] == "{%p for item in items %}"

    # tokens 正确标注：表格 row3 单元格里 {{item.缺陷等级}}
    row3 = next(r for r in t0["rows"] if r["id"] == "t0r3")
    cell1 = row3["cells"][1]["blocks"][0]
    assert cell1["text"] == "{{item.缺陷等级}}"
    kinds = {tok["kind"] for tok in cell1["tokens"]}
    assert "var" in kinds
    # tokens 区间应落在标签上
    for tok in cell1["tokens"]:
        assert cell1["text"][tok["start"]:tok["end"]] == tok["text"]


def test_parse_stable(migrated_template_bytes):
    """同一份 docx 两次 parse，块 id 与文本完全一致。"""
    m1 = parse_docx(migrated_template_bytes)
    m2 = parse_docx(migrated_template_bytes)
    assert m1 == m2


# --------------------------------------------------------------------------- #
# 2. 原样回写
# --------------------------------------------------------------------------- #
def test_apply_noop_preserves_xml(migrated_template_bytes):
    model = parse_docx(migrated_template_bytes)
    out = apply_model(migrated_template_bytes, model)

    # python-docx 能打开
    doc = Document(BytesIO(out))
    assert len(doc.tables) == 1

    # 未编辑段落 p1 的 XML 与原文件逐段一致
    def para_xml(b, idx):
        z = zipfile.ZipFile(BytesIO(b))
        root = etree.fromstring(z.read("word/document.xml"))
        body = root.find(qn("body"))
        ps = [ch for ch in body if etree.QName(ch).localname == "p"]
        return etree.tostring(ps[idx])

    assert para_xml(migrated_template_bytes, 1) == para_xml(out, 1)

    # styles.xml / numbering.xml 字节不变
    zo = zipfile.ZipFile(BytesIO(migrated_template_bytes))
    zn = zipfile.ZipFile(BytesIO(out))
    for part in ("word/styles.xml", "word/numbering.xml", "word/settings.xml",
                 "word/fontTable.xml"):
        assert zo.read(part) == zn.read(part), part + " 被改动"


def test_apply_noop_still_renderable(migrated_template_bytes, sample_data):
    """原样 apply 后，用示例数据 docxtpl 渲染仍产出 3 张表（与原模板一致）。"""
    model = parse_docx(migrated_template_bytes)
    out = apply_model(migrated_template_bytes, model)
    tpl = DocxTemplate(BytesIO(out))
    tpl.render(sample_data)
    buf = BytesIO(); tpl.save(buf)
    z = zipfile.ZipFile(BytesIO(buf.getvalue()))
    root = etree.fromstring(z.read("word/document.xml"))
    body = root.find(qn("body"))
    assert len(body.findall(qn("tbl"))) == len(sample_data["items"]) == 3


# --------------------------------------------------------------------------- #
# 3. 改文字保样式
# --------------------------------------------------------------------------- #
def test_apply_change_text_preserves_style(migrated_template_bytes):
    model = parse_docx(migrated_template_bytes)
    # 改 p0 控制段文本
    for b in model["blocks"]:
        if b["id"] == "p0":
            b["text"] = "{%p for vuln in items %}"
    out = apply_model(migrated_template_bytes, model)

    # 生效
    z = zipfile.ZipFile(BytesIO(out))
    root = etree.fromstring(z.read("word/document.xml"))
    body = root.find(qn("body"))
    ps = [ch for ch in body if etree.QName(ch).localname == "p"]
    assert "".join(t.text or "" for t in ps[0].iter(qn("t"))) == "{%p for vuln in items %}"

    # p1（未改）XML 逐段一致
    def para_xml(b, idx):
        zz = zipfile.ZipFile(BytesIO(b))
        rr = etree.fromstring(zz.read("word/document.xml"))
        bb = rr.find(qn("body"))
        pps = [ch for ch in bb if etree.QName(ch).localname == "p"]
        return etree.tostring(pps[idx])
    assert para_xml(migrated_template_bytes, 1) == para_xml(out, 1)


def test_apply_chinese_font_eastasia_kept():
    """改中文段落后，首个 run 的 eastAsia 字体 rPr 被保留。"""
    doc = Document()
    p = doc.add_paragraph()
    r = p.add_run("旧文本")
    rpr = r._element.get_or_add_rPr()
    rfonts = rpr.makeelement(qn("rFonts"), {qn("eastAsia"): "宋体"})
    rpr.insert(0, rfonts)
    buf = BytesIO(); doc.save(buf); raw = buf.getvalue()

    model = parse_docx(raw)
    model["blocks"][0]["text"] = "新中文文本"
    out = apply_model(raw, model)

    z = zipfile.ZipFile(BytesIO(out))
    root = etree.fromstring(z.read("word/document.xml"))
    east = root.find(".//" + qn("rFonts"))
    assert east.get(qn("eastAsia")) == "宋体"
    assert "新中文文本" in "".join(root.itertext())


# --------------------------------------------------------------------------- #
# 4. 占位符插入/删除 + 渲染
# --------------------------------------------------------------------------- #
def test_apply_insert_scalar_and_render(migrated_template_bytes, sample_data):
    model = parse_docx(migrated_template_bytes)
    # 在 p0 前插入一个新段落，含新标量 {{新变量}}
    new_block = {"id": "new-p0", "type": "paragraph", "text": "补充说明：{{新变量}}"}
    model["blocks"].insert(0, new_block)
    out = apply_model(migrated_template_bytes, model)

    # 渲染时提供新变量
    data = dict(sample_data)
    data["新变量"] = "插入的标量值"
    tpl = DocxTemplate(BytesIO(out))
    tpl.render(data)
    buf = BytesIO(); tpl.save(buf)
    z = zipfile.ZipFile(BytesIO(buf.getvalue()))
    text = "".join(etree.fromstring(z.read("word/document.xml")).itertext())
    assert "插入的标量值" in text
    assert "{{新变量}}" not in text


def test_apply_unbalanced_tag_raises(migrated_template_bytes):
    model = parse_docx(migrated_template_bytes)
    for b in model["blocks"]:
        if b["id"] == "p2":  # {%p if not loop.last %}
            b["text"] = "{%p if show"  # 未闭合
    out = apply_model(migrated_template_bytes, model)
    with pytest.raises(Exception) as ei:
        tpl = DocxTemplate(BytesIO(out))
        tpl.render({})
    assert "end" in str(ei.value).lower() or "if" in str(ei.value).lower()


# --------------------------------------------------------------------------- #
# 5. 表格增删行 / 改单元格 / 循环行
# --------------------------------------------------------------------------- #
def _make_table_doc():
    doc = Document()
    doc.add_paragraph("表头说明")
    t = doc.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text = "名称"
    t.rows[0].cells[1].text = "等级"
    t.rows[1].cells[0].text = "{{item.name}}"
    t.rows[1].cells[1].text = "{{item.grade}}"
    buf = BytesIO(); doc.save(buf); return buf.getvalue()


def test_apply_table_loop_row_renders_count():
    raw = _make_table_doc()
    model = parse_docx(raw)
    for b in model["blocks"]:
        if b["type"] == "table":
            for row in b["rows"]:
                if row["id"] == "t0r1":
                    row["loop_row"] = True
                    row["loop"] = {"var": "item", "iterable": "items"}
    out = apply_model(raw, model)

    tpl = DocxTemplate(BytesIO(out))
    tpl.render({"items": [{"name": "A", "grade": "高"},
                          {"name": "B", "grade": "中"},
                          {"name": "C", "grade": "低"}]})
    buf = BytesIO(); tpl.save(buf)
    z = zipfile.ZipFile(BytesIO(buf.getvalue()))
    root = etree.fromstring(z.read("word/document.xml"))
    rows = root.find(".//" + qn("tbl")).findall(qn("tr"))
    # 表头 1 行 + 数据 3 行
    assert len(rows) == 4


def test_apply_table_add_delete_edit_row():
    raw = _make_table_doc()
    model = parse_docx(raw)
    for b in model["blocks"]:
        if b["type"] == "table":
            rows = b["rows"]
            # 改表头单元格
            rows[0]["cells"][0]["blocks"][0]["text"] = "缺陷名称"
            # 删数据行 t0r1
            b["rows"] = [rows[0]]
    out = apply_model(raw, model)
    z = zipfile.ZipFile(BytesIO(out))
    root = etree.fromstring(z.read("word/document.xml"))
    rows = root.find(".//" + qn("tbl")).findall(qn("tr"))
    # 只剩表头 1 行
    assert len(rows) == 1
    assert "缺陷名称" in "".join(rows[0].itertext())


# --------------------------------------------------------------------------- #
# 6. unsupported 文本框原样保留
# --------------------------------------------------------------------------- #
def _make_docx_with_textbox():
    doc = Document()
    doc.add_paragraph("正常段落")
    # 手工注入 w:txbxContent 到第二个段落
    p = doc.add_paragraph()
    r = p.add_run()
    txbx = r._element.makeelement(qn("pict"), {})
    txbxContent = txbx.makeelement(qn("txbxContent"), {})
    inner_p = txbxContent.makeelement(qn("p"), {})
    inner_r = inner_p.makeelement(qn("r"), {})
    inner_t = inner_r.makeelement(qn("t"), {})
    inner_t.text = "我是文本框"
    inner_r.append(inner_t); inner_p.append(inner_r); txbxContent.append(inner_p)
    txbx.append(txbxContent); r._element.append(txbx)
    buf = BytesIO(); doc.save(buf); return buf.getvalue()


def test_parse_and_apply_textbox_unsupported():
    raw = _make_docx_with_textbox()
    model = parse_docx(raw)
    kinds = [b["type"] for b in model["blocks"]]
    assert "unsupported" in kinds
    # 文本框段落在模型里是 unsupported
    uns = [b for b in model["blocks"] if b["type"] == "unsupported"]
    assert uns and "文本框" in uns[0]["note"]

    # 即使模型里"忘了"回传 unsupported 块，apply 也必须原样保留它的 XML
    stripped = {"blocks": [b for b in model["blocks"] if b["type"] != "unsupported"]}
    out = apply_model(raw, stripped)
    zn = zipfile.ZipFile(BytesIO(out))
    out_xml = zn.read("word/document.xml")
    # txbxContent 与文本框文字必须原样保留
    assert out_xml.find(b"txbxContent") != -1
    assert out_xml.find("我是文本框".encode("utf-8")) != -1
    # 正常段落仍在
    assert out_xml.find("正常段落".encode("utf-8")) != -1


# --------------------------------------------------------------------------- #
# 7. 安全：docm/超大文件拦截
# --------------------------------------------------------------------------- #
def test_security_rejects_docm(blank_template_bytes, tmp_path):
    # 把 blank-template 包装成 docm：篡改 Content_Types 加宏标识
    z = zipfile.ZipFile(BytesIO(blank_template_bytes))
    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
        for item in z.infolist():
            data = z.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"wordprocessingml.document.main+xml",
                                    b"wordprocessingml.document.macroEnabled.main+xml")
            zo.writestr(item, data)
    docm_bytes = out.getvalue()
    with pytest.raises(SecurityError) as ei:
        validate_docx(docm_bytes)
    assert ei.value.status == 415


def test_security_rejects_oversize():
    import app.config as cfg
    big = b"PK\x03\x04" + b"\x00" * (cfg.MAX_TEMPLATE_BYTES + 10)
    with pytest.raises(SecurityError) as ei:
        validate_docx(big)
    assert ei.value.status == 413


def test_apply_endpoint_rejects_docm(migrated_template_bytes):
    """走 HTTP 路由：docm 上传 /apply 仍被拒。"""
    z = zipfile.ZipFile(BytesIO(migrated_template_bytes))
    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
        for item in z.infolist():
            data = z.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"wordprocessingml.document.main+xml",
                                    b"wordprocessingml.document.macroEnabled.main+xml")
            zo.writestr(item, data)
    resp = client.post(
        "/apply",
        files={"file": ("t.docm", out.getvalue(), DOCX_CT)},
        data={"model": json.dumps({"blocks": []})},
    )
    assert resp.status_code == 415
