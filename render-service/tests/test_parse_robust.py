# -*- coding: utf-8 -*-
"""parse_docx 的健壮性：坏元素不连累整份文档 + 不可编辑元素带出原文。

背景（用户实报）：「该模板含暂不支持的元素，已以只读模式打开。请在 Word 中处理后
重新上传。应该要所有元素都支持啊」。追下去发现是两个问题叠在一起：

  1. 文案说"部分元素不支持"，实际撞到的是**整份 parse 抛异常**的兜底分支 ——
     一个元素解析失败，整份模板就打不开，用户被引导去 Word 里反复删东西；
  2. 即使真的只是局部降级，前端也只显示一个空白方块 +「需在 Word 处理」，
     用户既看不到盒子里写了什么，也不知道到底该改什么。

这份测试把三件事钉住：
  A. 不可编辑元素（文本框/形状/浮动锚定）的块必须**带出盒内原文**（text 字段）；
  B. 某个元素解析抛异常时**只降级它自己**，前后段落全保留，id 仍按 body 顺序稳定；
  C. 同一份 docx 两次 parse 结果逐字节一致（docapply 的 id 对齐依赖这个前提）。

可直接 `python tests/test_parse_robust.py` 跑，也兼容 pytest。
"""
import json
import os
import sys
import zipfile
from io import BytesIO

from docx import Document

RENDER_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if RENDER_DIR not in sys.path:
    sys.path.insert(0, RENDER_DIR)

from app import docmodel  # noqa: E402
from app.docmodel import parse_docx  # noqa: E402
from app.preprocess import qn  # noqa: E402


# --------------------------------------------------------------------------- #
# 构造样本 docx
# --------------------------------------------------------------------------- #
def _docx_with_textbox(inner_text: str) -> bytes:
    """一个普通段落 + 一个装文本框的段落 + 一个普通段落。"""
    doc = Document()
    doc.add_paragraph("前一段")
    p = doc.add_paragraph()
    r = p.add_run()
    pict = r._element.makeelement(qn("pict"), {})
    txbx_content = pict.makeelement(qn("txbxContent"), {})
    inner_p = txbx_content.makeelement(qn("p"), {})
    inner_r = inner_p.makeelement(qn("r"), {})
    inner_t = inner_r.makeelement(qn("t"), {})
    inner_t.text = inner_text
    inner_r.append(inner_t)
    inner_p.append(inner_r)
    txbx_content.append(inner_p)
    pict.append(txbx_content)
    r._element.append(pict)
    doc.add_paragraph("后一段")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _docx_with_table() -> bytes:
    """段落 / 表格 / 段落，用来验证"表格炸了不影响前后段落"。"""
    doc = Document()
    doc.add_paragraph("表格前")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "A"
    t.cell(0, 1).text = "B"
    t.cell(1, 0).text = "C"
    t.cell(1, 1).text = "D"
    doc.add_paragraph("表格后")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# A. 不可编辑元素带出盒内原文
# --------------------------------------------------------------------------- #
def test_textbox_block_carries_inner_text():
    model = parse_docx(_docx_with_textbox("我是文本框里的字"))
    uns = [b for b in model["blocks"] if b["type"] == "unsupported"]
    assert len(uns) == 1, model["blocks"]
    # note 保留"文本框"字样（旧测试与前端标签都依赖它）
    assert "文本框" in uns[0]["note"]
    # 关键：盒内原文必须带出来，否则前端只能显示一个空白方块
    assert uns[0]["text"] == "我是文本框里的字"


def test_textbox_note_does_not_tell_user_to_edit_in_word():
    """文案不能再说"需在 Word 处理" —— 生成文档时这些元素原样保留，根本不用改。"""
    model = parse_docx(_docx_with_textbox("x"))
    uns = [b for b in model["blocks"] if b["type"] == "unsupported"]
    assert "需在 Word 处理" not in uns[0]["note"]
    assert "原样保留" in uns[0]["note"]


def test_textbox_does_not_affect_neighbouring_paragraphs():
    model = parse_docx(_docx_with_textbox("盒内"))
    texts = [b.get("text") for b in model["blocks"]]
    assert texts == ["前一段", "盒内", "后一段"], model["blocks"]


def test_warnings_are_deduped():
    """同类降级出现多次时，告警只留一条，别把有用的信息淹掉。"""
    doc = Document()
    for i in range(5):
        p = doc.add_paragraph()
        r = p.add_run()
        pict = r._element.makeelement(qn("pict"), {})
        content = pict.makeelement(qn("txbxContent"), {})
        ip = content.makeelement(qn("p"), {})
        ir = ip.makeelement(qn("r"), {})
        it = ir.makeelement(qn("t"), {})
        it.text = "框%d" % i
        ir.append(it)
        ip.append(ir)
        content.append(ip)
        pict.append(content)
        r._element.append(pict)
    buf = BytesIO()
    doc.save(buf)
    model = parse_docx(buf.getvalue())
    assert len([b for b in model["blocks"] if b["type"] == "unsupported"]) == 5
    box_warnings = [w for w in model["warnings"] if "文本框" in w]
    assert len(box_warnings) == 1, model["warnings"]


# --------------------------------------------------------------------------- #
# B. 单个元素解析失败只降级它自己
# --------------------------------------------------------------------------- #
def test_one_broken_table_degrades_only_itself(monkeypatch=None):
    raw = _docx_with_table()
    boom = RuntimeError("模拟结构异常")
    orig = docmodel._parse_table

    def _raise(*a, **kw):
        raise boom

    docmodel._parse_table = _raise
    try:
        model = parse_docx(raw)
    finally:
        docmodel._parse_table = orig

    kinds = [b["type"] for b in model["blocks"]]
    texts = [b.get("text") for b in model["blocks"]]
    # 整份文档必须解析成功，前后段落都在
    assert kinds == ["paragraph", "unsupported", "paragraph"], model["blocks"]
    assert texts[0] == "表格前"
    assert texts[2] == "表格后"
    # 降级块的 id 必须还在表格的位置上（t0），保证 apply 侧能对齐
    assert model["blocks"][1]["id"] == "t0"
    assert "表格结构异常" in model["blocks"][1]["note"]
    assert "RuntimeError" in model["blocks"][1]["note"]


def test_broken_element_keeps_id_sequence_contiguous():
    """降级后 id 不能出现空洞或重号，否则 apply 会错位。"""
    raw = _docx_with_table()
    orig = docmodel._parse_table

    def _raise(*a, **kw):
        raise ValueError("boom")

    docmodel._parse_table = _raise
    try:
        model = parse_docx(raw)
    finally:
        docmodel._parse_table = orig

    ids = [b["id"] for b in model["blocks"]]
    assert ids == ["p0", "t0", "p1"], ids


# --------------------------------------------------------------------------- #
# C. 解析结果稳定（apply 的 id 对齐前提）
# --------------------------------------------------------------------------- #
def test_parse_is_deterministic():
    raw = _docx_with_textbox("稳定的")
    a = json.dumps(parse_docx(raw), ensure_ascii=False, sort_keys=True)
    b = json.dumps(parse_docx(raw), ensure_ascii=False, sort_keys=True)
    assert a == b


def test_sample_template_still_parses_clean():
    """真实审计报告模板必须零告警解析 —— 防止降级逻辑误伤正常文档。"""
    path = os.path.join(os.path.dirname(RENDER_DIR), "samples", "templates", "audit-report.docx")
    if not os.path.exists(path):
        return  # 样例缺失时不阻断
    with open(path, "rb") as f:
        model = parse_docx(f.read())
    assert model["blocks"], "样例模板解析出 0 个块"
    assert model["warnings"] == [], model["warnings"]
    kinds = {b["type"] for b in model["blocks"]}
    assert kinds <= {"heading", "paragraph", "list", "table", "placeholder", "unsupported"}


# --------------------------------------------------------------------------- #
# D. 页眉/页脚占位符：不在画布内，但必须提示
# --------------------------------------------------------------------------- #
def test_header_footer_tokens_are_announced():
    """页眉里的占位符不在编辑画布内 —— 不提示的话用户会以为模板里没这个变量。"""
    doc = Document()
    doc.add_paragraph("正文")
    doc.sections[0].header.paragraphs[0].text = "编号：{{报告编号}}"
    doc.sections[0].footer.paragraphs[0].text = "{{单位名称}}"
    buf = BytesIO()
    doc.save(buf)
    model = parse_docx(buf.getvalue())
    joined = "；".join(model["warnings"])
    assert "页眉" in joined and "报告编号" in joined, model["warnings"]
    assert "页脚" in joined and "单位名称" in joined, model["warnings"]


def test_no_header_footer_warning_when_no_tokens():
    """页眉页脚只有静态文字时不该产生任何提示，否则正常模板会被噪声淹没。"""
    doc = Document()
    doc.add_paragraph("正文")
    doc.sections[0].header.paragraphs[0].text = "内部资料 注意保密"
    doc.sections[0].footer.paragraphs[0].text = "第 1 页"
    buf = BytesIO()
    doc.save(buf)
    model = parse_docx(buf.getvalue())
    assert model["warnings"] == [], model["warnings"]


if __name__ == "__main__":
    ok = fail = 0
    for name, fn in sorted(globals().items()):
        if not name.startswith("test_") or not callable(fn):
            continue
        try:
            fn()
            print("PASS  %s" % name)
            ok += 1
        except AssertionError as e:
            print("FAIL  %s -> %s" % (name, e))
            fail += 1
        except Exception as e:  # noqa: BLE001
            print("ERROR %s -> %s: %s" % (name, type(e).__name__, e))
            fail += 1
    print("\n%d passed, %d failed" % (ok, fail))
    sys.exit(1 if fail else 0)
