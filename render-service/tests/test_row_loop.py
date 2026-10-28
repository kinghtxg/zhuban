# -*- coding: utf-8 -*-
"""表格行循环：{%tr for %} 的两种写法都必须真的循环。

背景：docxtpl 0.20.2 自带的 {%tr %} 处理有缺陷 —— 它的正则前缀是贪婪的，
回溯时取到的是同一行里**最后**一个 {%tr 标签。手册教的写法（首格 for、
末格 endfor、同一行）因此把整行换成 `{% endfor %}`，`for` 消失，报错是
「Encountered unknown tag 'endfor'」，完全看不出跟表格有关。

我们在 preprocess 里自己把行级标签搬到行外，所以这里钉住三件事：
循环真的展开、空数据不产生空行、两种写法结果一致。
"""
import zipfile
from io import BytesIO

from docx import Document
from lxml import etree

from app.preprocess import normalize_row_tags_docx, preprocess_docx
from app.scanner import scan_docx
from app.render import render_template

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
HEAD = ["名称", "数值", "环比"]


def _rows(docx_bytes):
    root = etree.fromstring(zipfile.ZipFile(BytesIO(docx_bytes)).read("word/document.xml"))
    out = []
    for tbl in root.iter(W + "tbl"):
        for tr in tbl.findall(W + "tr"):
            out.append(["".join(c.itertext()).strip() for c in tr.findall(W + "tc")])
    return out


def _same_row(items):
    """手册约定：for 在首格开头、endfor 在末格末尾，同一行。"""
    d = Document()
    d.add_paragraph("指标表")
    t = d.add_table(rows=2, cols=3)
    for i, h in enumerate(HEAD):
        t.rows[0].cells[i].text = h
    r = t.rows[1]
    r.cells[0].text = "{%tr for m in 指标 %}{{ m.名称 }}"
    r.cells[1].text = "{{ m.数值 }}"
    r.cells[2].text = "{{ m.环比 }}{%tr endfor %}"
    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()


def _tag_own_row(items):
    """标签各占一整行的写法。"""
    d = Document()
    d.add_paragraph("指标表")
    t = d.add_table(rows=4, cols=3)
    for i, h in enumerate(HEAD):
        t.rows[0].cells[i].text = h
    t.rows[1].cells[0].text = "{%tr for m in 指标 %}"
    t.rows[2].cells[0].text = "{{ m.名称 }}"
    t.rows[2].cells[1].text = "{{ m.数值 }}"
    t.rows[2].cells[2].text = "{{ m.环比 }}"
    t.rows[3].cells[0].text = "{%tr endfor %}"
    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()


DATA = {"指标": [
    {"名称": "A1", "数值": "1", "环比": "+1%"},
    {"名称": "B1", "数值": "2", "环比": "+2%"},
]}


def test_row_loop_same_row_renders():
    rows = _rows(render_template(_same_row(None), DATA))
    assert rows == [HEAD, ["A1", "1", "+1%"], ["B1", "2", "+2%"]], rows


def test_row_loop_tag_own_row_renders_without_blank_rows():
    rows = _rows(render_template(_tag_own_row(None), DATA))
    assert rows == [HEAD, ["A1", "1", "+1%"], ["B1", "2", "+2%"]], rows
    assert [] not in rows, "标签行必须整行删掉，不能变成空行留在成品里"


def test_row_loop_empty_data_keeps_only_header():
    rows = _rows(render_template(_same_row(None), {"指标": []}))
    assert rows == [HEAD], rows


def test_row_loop_no_leftover_tags():
    out = render_template(_same_row(None), DATA)
    xml = zipfile.ZipFile(BytesIO(out)).read("word/document.xml").decode("utf-8")
    assert "{%" not in xml and "{{" not in xml
    assert "<w:tr" in xml


def test_preprocess_moves_row_tags_out_of_row():
    """规范化后的 XML 里不能再出现 {%tr %}，否则 docxtpl 那条正则会再次误伤。"""
    fixed, _ = normalize_row_tags_docx(_same_row(None))
    xml = zipfile.ZipFile(BytesIO(fixed)).read("word/document.xml").decode("utf-8")
    assert "{%tr" not in xml
    assert "{% for m in 指标 %}" in xml and "{% endfor %}" in xml
    # 循环边界必须夹住同一个 <w:tr>，而不是把整张表都圈进去
    i_for = xml.index("{% for m in 指标 %}")
    i_row = xml.index("<w:tr", i_for)
    i_end = xml.index("{% endfor %}")
    assert i_for < i_row < i_end
    assert "</w:tr>" in xml[i_row:i_end]


def test_scan_preprocess_keeps_row_tags_visible():
    """扫描用的 preprocess_docx **不能**搬标签。

    搬走之后标签不再是 <w:t> 里的文字，扫描器读不到 `{%tr for %}`，就会把
    循环变量当成普通标量：示例 JSON 多出 "m"/"it" 两个键，投递时自己的
    校验报「字段 m 缺失」—— 模板明明是对的，用户却怎么也提交不过。
    """
    fixed, _, _ = preprocess_docx(_same_row(None))
    xml = zipfile.ZipFile(BytesIO(fixed)).read("word/document.xml").decode("utf-8")
    assert "{%tr for m in 指标 %}" in xml

    scan = scan_docx(_same_row(None))
    names = [v["name"] for v in scan["variables"]]
    assert "指标" in names and "m" not in names, names
    assert scan["json_schema"]["required"] == ["指标"], scan["json_schema"]["required"]


def test_row_loop_nested_in_paragraph_loop():
    """行循环套在整段循环里（每组分一张表）：两组各出 2 行，表头各一份。"""
    d = Document()
    d.add_paragraph("{%p for g in 分组 %}")
    d.add_paragraph("{{ g.分组 }}")
    t = d.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text = "问题"
    t.rows[0].cells[1].text = "等级"
    t.rows[1].cells[0].text = "{%tr for it in g.条目 %}{{ it.问题 }}"
    t.rows[1].cells[1].text = "{{ it.等级 }}{%tr endfor %}"
    d.add_paragraph("{%p endfor %}")
    buf = BytesIO()
    d.save(buf)

    data = {"分组": [
        {"分组": "账号", "条目": [{"问题": "弱口令", "等级": "高"}]},
        {"分组": "备份", "条目": [{"问题": "无演练", "等级": "中"}, {"问题": "日志短", "等级": "低"}]},
    ]}
    rows = _rows(render_template(buf.getvalue(), data))
    assert rows == [
        ["问题", "等级"], ["弱口令", "高"],
        ["问题", "等级"], ["无演练", "中"], ["日志短", "低"],
    ], rows
